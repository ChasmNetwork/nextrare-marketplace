# python3 scripts/deck-pptx.py → docs/NextRare-Deck.pptx  (native text boxes, imports cleanly into Canva / Keynote / Slides)
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BG, INK, MUTED, ACC, RED, TILE = "0B0B0D", "F4F4F5", "A1A1AA", "FF734C", "DD2023", "17171B"
LOGO = "public/brand/nextrare-logo.png"
APP = "nextrare-marketplace.vercel.app"
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def text(slide, x, y, w, h, runs, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: str | list of paragraphs; each paragraph str | list of (text, {bold, color}) tuples"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"): setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align; p.line_spacing = spacing
        if i: p.space_before = Pt(size * 0.45)
        for seg in (para if isinstance(para, list) else [(para, {})]):
            t, o = (seg, {}) if isinstance(seg, str) else seg
            r = p.add_run(); r.text = t; f = r.font; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color)); f.name = "Sora"
    return tb

def tile(slide, x, y, w, h, big, small):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.08; s.fill.solid(); s.fill.fore_color.rgb = rgb(TILE); s.line.color.rgb = rgb("2A2A31"); s.line.width = Pt(0.75)
    text(slide, x + 0.22, y + 0.18, w - 0.44, 0.6, big, size=30, bold=True)
    text(slide, x + 0.22, y + 0.85, w - 0.44, h - 0.95, small, size=11, color=MUTED)

def slide(n, kicker, head, accent, body_fn):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.5), Inches(12.133), Inches(6.5))
    box.adjustments[0] = 0.03; box.fill.solid(); box.fill.fore_color.rgb = rgb("0E0E11"); box.line.color.rgb = rgb(RED); box.line.width = Pt(2)
    if n > 1: s.shapes.add_picture(LOGO, Inches(10.3), Inches(0.88), height=Inches(0.36))
    text(s, 1.1, 0.95, 11, 0.35, f"{n:02d} · {kicker.upper()}", size=11, color=MUTED)
    text(s, 1.1, 1.3, 11.2, 1.1, [[(head + " " if head else "", {}), (accent, {"color": ACC})]], size=40, bold=True)
    body_fn(s)
    text(s, 0.6, 7.05, 6, 0.3, "NextRare", size=10, color="71717A")
    text(s, 6.7, 7.05, 6.03, 0.3, f"{n} / 10", size=10, color="71717A", align=PP_ALIGN.RIGHT)

def bullets(s, items, y=2.55, size=19, h=4.2):
    text(s, 1.1, y, 11.1, h, [[("•  ", {"color": MUTED})] + (it if isinstance(it, list) else [(it, {})]) for it in items], size=size, spacing=1.2)

def numbered(s, items, y=2.55, size=19):
    text(s, 1.1, y, 11.1, 4.2, [[(f"{i+1}.  ", {"color": ACC, "bold": True})] + (it if isinstance(it, list) else [(it, {})]) for i, it in enumerate(items)], size=size, spacing=1.2)

B = {"bold": True}

# 01
def s1(s):
    s.shapes.add_picture(LOGO, Inches(1.1), Inches(1.35), height=Inches(0.9))
    text(s, 1.1, 2.5, 11, 1.0, [[("A trading card show ", {}), ("at your fingertips.", {"color": ACC})]], size=40, bold=True)
    text(s, 1.1, 3.75, 11, 0.6, "Marketplace and gacha in one. List a card once, sell it two ways, earn while it waits.", size=17, color=MUTED)
    text(s, 1.1, 4.6, 11, 0.5, "KC Thee, Li Ho  ·  live on Solana devnet  ·  built at Startup Village Borneo, Sept 6–8 2026", size=14, color=MUTED)
    text(s, 1.1, 5.1, 11, 0.5, APP, size=14, color=ACC)
slide(1, "Title", "", "", s1)

slide(2, "Problem", "List. Wait.", "Hope.", lambda s: bullets(s, [
    "Card demand never switches off, it moves between formats: Web2 marketplaces, tokenised cards, onchain gacha, thousands of local shops. Nothing connects them.",
    "A collector who lists a graded card at a fair price waits weeks. The only fast exit is selling under market.",
    "While it waits, the card earns nothing."]))

slide(3, "Solution", "List once.", "Sell two ways.", lambda s: bullets(s, [
    "Your card is buyable at your price, and sits inside a gacha pack at the same time.",
    "A ripper keeps it: you get full price, instantly. A ripper takes cash instead: your card stays, and you get paid rent for waiting.",
    "The card never leaves your wallet. A Metaplex Core freeze plugin holds it, not us.",
    "This is the onchain version of the stake-to-earn layer inside NextRare, the gacha we already run."]))

def s4(s):
    numbered(s, ["Connect Phantom on devnet", "Tap “Get free cards”: two slabs land in your wallet", "List one. It is on sale and in a pack at once",
                 "Rip a pack. Keep the card, or take 85% cash now", "Watch rent land in “My cards & earnings”"])
    text(s, 1.1, 6.1, 11, 0.4, f"{APP}   ·   {APP}/traction", size=14, color=ACC)
slide(4, "Demo", "Sixty seconds,", "no signup.", s4)

def s5(s):
    text(s, 1.1, 2.45, 11, 0.35, "NextRare gacha, live since January 2026  ·  8 months  ·  no paid acquisition", size=13, color=MUTED)
    tiles = [("$1.26M", "gross merchandise value"), ("$836k", "transaction volume · 8,074 orders"), ("256", "paying users · 2,537 signed in"),
             ("$3,265", "volume per paying user"), ("$427k", "cash returned on sell-backs"), ("1,046", "packs opened in the offline pilot since May")]
    for i, (b, sm) in enumerate(tiles): tile(s, 1.1 + (i % 3) * 3.75, 2.9 + (i // 3) * 1.5, 3.6, 1.35, b, sm)
    bullets(s, [[("A typical cycle: pull a card worth ", {}), ("$42", B), (", sell it back for ", {}), ("$38", B), (". ", {}), ("86%", B), (" of cards are sold back, so inventory recycles instead of shipping out.", {})],
                [("This marketplace: ", {}), ("2%", B), (" on direct buys. On a cash-out the ripper gets ", {}), ("85%", B), (", the 15% gap is profit, ", {}), ("split 50/50", B), (" with the sellers in that pack.", {})]],
            y=5.95, size=13, h=1.0)
slide(5, "Traction & business model", "The gacha already", "works.", s5)

slide(6, "Market · why now · why Solana", "Validated in the West.", "Empty in SEA.", lambda s: bullets(s, [
    "Collector Crypt did ~$153M of gacha spend in Q1 2026. Courtyard went from $50k to ~$50M a month. Beezie has $170M+ cumulative GMV. All Western, online only, single brand.",
    "Southeast Asia is the fastest-growing TCG market, deep collector culture, dense card-shop networks, and no aggregated hybrid platform.",
    "Why Solana: non-custodial listings via Core plugins, one-transaction settlement, fees in cents so a $5 pack works, and every roll provable from the payment transaction."], size=18))

slide(7, "Competition", "One listing,", "both exits.", lambda s: bullets(s, [
    [("Collector Crypt · Courtyard · Beezie · OpenGacha", B), (": buy their own players, hold their own inventory, one exit per listing, sellers earn nothing while they wait.", {})],
    [("Tensor · Magic Eden", B), (": fixed price only, idle listings earn nothing.", {})],
    [("Us", B), (": both exits on one listing, non-custodial, the waiting itself pays, and physical rails in SEA.", {})]]))

slide(8, "GTM · 3-month plan", "Distribution we", "already have.", lambda s: numbered(s, [
    [("Activate", B), (" the 2,537 registered NextRare users against 256 paying: they become the first sellers and rippers on mainnet.", {})],
    [("Penang flagship", B), (" opens October 2026 with $300k+ of consignor inventory. The app is the shop’s wallet, so every purchase onboards a user.", {})],
    [("Partners", B), (": Speculate and CatchaCard live as white-label channels; Slabz and Discover Collectibles confirmed. Their inventory lists here, their users rip here.", {})]], size=18))

slide(9, "Team", "Why", "us.", lambda s: bullets(s, [
    [("KC Thee", B), (", CEO: helped scale Binance across Southeast Asia, P2P and futures. Advisor to Virtuals.", {})],
    [("John Koh", B), (", CTO: former ML engineer at Eligible (YC 2012), Bitcoin since 2011.", {})],
    [("Huey Lau", B), (", COO: built Discover Collectibles, 130+ merchants across SEA retail.", {})],
    [("Li Ho", B), (", engineering: built the NextRare app and shipped this marketplace in three days.", {})]]))

def s10(s):
    for i, (b, sm) in enumerate([("$1M", "raise"), ("$10M", "post-money valuation"), ("24 mo", "runway")]): tile(s, 1.1 + i * 3.75, 2.55, 3.6, 1.35, b, sm)
    bullets(s, [[("45%", B), (" user acquisition: activate the 2,537 registered users, partner and merchant onboarding.", {})],
                [("25%", B), (" development: one more engineer, mainnet launch of this marketplace, membership out of beta.", {})],
                [("30%", B), (" operations: licensing, on/off-ramp, vaulting and authentication.", {})],
                [("Plus intros: Solana Foundation, Metaplex, Superteam MY collectors and card shops.  Try it now: ", {}), (APP, {"color": ACC})]], y=4.2, size=16, h=2.6)
slide(10, "The ask", "Raising $1M at", "$10M post-money.", s10)

out = "docs/NextRare-Deck.pptx"; prs.save(out); print("wrote", out)
# sanity: reopen and count
p2 = Presentation(out); print("slides", len(p2.slides), "shapes/slide", [len(sl.shapes) for sl in p2.slides])
